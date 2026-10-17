#!/usr/bin/env python3
"""Clone XLSX packages, edit scalar cells, refresh formula caches and render copies.

openpyxl supplies workbook/cell inspection. Surgical OOXML edits avoid its
lossy load/save round trip for charts, drawings and extension parts.
"""
import argparse
from datetime import date, datetime, time, timedelta
import hashlib
import importlib.util
from html import escape
import json
import math
import os
from pathlib import Path
import posixpath
import re
import shutil
import signal
import socket
import subprocess
import sys
import tempfile
import warnings
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.cell.rich_text import CellRichText
from openpyxl.utils.cell import coordinate_to_tuple
from openpyxl.utils.datetime import to_excel, time_to_days, timedelta_to_days

MAIN = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
CELL = re.compile(rb'<c\b[^>]*?\br="([^"]+)"[^>]*?(?:/>|>.*?</c>)', re.S)
BLOCKED = ('vbaProject', 'externalLinks/', 'connections.xml', 'queryTables/',
           'activeX/', 'embeddings/', '_xmlsignatures/', 'xl/macrosheets/')


class WorkflowError(Exception):
    pass


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def excel_scalar(value, epoch):
    if isinstance(value, (date, datetime)):
        return to_excel(value, epoch)
    if isinstance(value, time):
        return time_to_days(value)
    if isinstance(value, timedelta):
        return timedelta_to_days(value)
    return value


def package(path):
    path = Path(path).resolve()
    if path.suffix.lower() != '.xlsx':
        raise WorkflowError('Only ordinary .xlsx workbooks are supported.')
    with ZipFile(path) as z:
        if len(z.namelist()) != len(set(z.namelist())):
            raise WorkflowError('Duplicate workbook ZIP members are unsupported.')
        if sum(i.file_size for i in z.infolist()) > 200_000_000:
            raise WorkflowError('Workbook exceeds the 200 MB expanded-package limit.')
        return {n: z.read(n) for n in z.namelist()}


def sheet_parts(parts):
    wb = ET.fromstring(parts['xl/workbook.xml'])
    rels = ET.fromstring(parts['xl/_rels/workbook.xml.rels'])
    targets = {r.attrib['Id']: r.attrib['Target'] for r in rels}
    result = {}
    for sheet in wb.find(f'{{{MAIN}}}sheets'):
        target = targets[sheet.attrib[f'{{{REL}}}id']]
        name = target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/' + target)
        result[sheet.attrib['name']] = name
    return result


def inspect(source):
    parts = package(source)
    blocked = [n for n in parts if any(token in n for token in BLOCKED)]
    if b'macroEnabled' in parts.get('[Content_Types].xml', b''):
        blocked.append('macro-enabled content type')
    sheets = sheet_parts(parts)
    for sheet, part in sheets.items():
        if not part.startswith('xl/worksheets/'):
            blocked.append('non-worksheet sheet: ' + sheet)
        for formula in ET.fromstring(parts[part]).iter(f'{{{MAIN}}}f'):
            if formula.get('t') in ('array', 'dataTable'):
                blocked.append('array/data-table formula in ' + sheet)
            text = formula.text or ''
            if re.search(r'(?i)\b(?:WEBSERVICE|DDE|RTD|STOCKHISTORY|IMAGE|IMPORTXML|IMPORTDATA|IMPORTRANGE|GOOGLEFINANCE)\s*\(', text) or re.search(r'\[[^]]+\][^!]*!', text):
                blocked.append('external-capable formula in ' + sheet)
    opaque = [n for n in parts if any(s in n for s in ('/charts/', '/drawings/', '/pivot', '/slicer', 'customXml/'))]
    return {'source_sha256': sha(source), 'sheets': list(sheets),
            'blocked_features': blocked, 'opaque_parts_preserved': opaque,
            'limitations': [
                'Scalar cells only; no row/column/sheet insertion or formatting redesign.',
                'Formula caches are refreshed; chart/pivot caches and objects are not rewritten.',
                'PDF/PNG previews reflect LibreOffice rendering, not pixel-identical Excel output.',
            ]}


def write_package(source, destination, replacements):
    destination = Path(destination)
    if destination.exists() or destination.is_symlink():
        raise WorkflowError('Output already exists; refusing to overwrite it.')
    created = False
    try:
        handle = destination.open('xb')
        created = True
        with handle, ZipFile(source) as src, ZipFile(handle, 'w') as dst:
            for info in src.infolist():
                dst.writestr(info, replacements.get(info.filename, src.read(info.filename)))
    except Exception:
        # Only remove a file created by this invocation, never a pre-existing path.
        if created and destination.exists():
            destination.unlink()
        raise


def cell_pattern(coordinate):
    return re.compile(rb'<c\b[^>]*?\br="' + re.escape(coordinate.encode()) + rb'"[^>]*?(?:/>|>.*?</c>)', re.S)


def content_cell(old, coordinate, value, formula=False, cached=False):
    if old:
        match = re.match(rb'<c\b([^>]*?)(/?)>', old)
        attrs = re.sub(rb'\s+t="[^"]*"', b'', match.group(1))
        inner = b'' if match.group(2) else old[match.end():-4]
    else:
        attrs, inner = (' r="' + coordinate + '"').encode(), b''
    for tag in (('v', 'is') if cached else ('f', 'v', 'is')):
        inner = re.sub(rb'<' + tag.encode() + rb'\b[^>]*?(?:/>|>.*?</' + tag.encode() + rb'>)', b'', inner, flags=re.S)
    kind, payload = '', ''
    if formula:
        payload = '<f>' + escape(value[1:]) + '</f><v></v>'
    elif value is None:
        payload = ''
    elif isinstance(value, bool):
        kind, payload = 'b', '<v>' + str(int(value)) + '</v>'
    elif isinstance(value, (int, float)):
        kind, payload = 'n', '<v>' + repr(value) + '</v>'
    elif cached:
        kind, payload = 'str', '<v>' + escape(value) + '</v>'
    else:
        kind = 'inlineStr'
        payload = '<is><t xml:space="preserve">' + escape(value) + '</t></is>'
    if kind:
        attrs += (' t="' + kind + '"').encode()
    # Cell content precedes extension data. Preserve all non-content children.
    if cached:
        formula_node = re.search(rb'<f\b[^>]*?(?:/>|>.*?</f>)', inner, re.S)
        if not formula_node:
            raise WorkflowError('Cache update requires an existing formula.')
        inner = inner[:formula_node.end()] + payload.encode() + inner[formula_node.end():]
    else:
        inner = payload.encode() + inner
    return b'<c' + attrs + b'>' + inner + b'</c>'


def replace_cell(xml, coordinate, value, formula=False, cached=False):
    if re.search(rb'<[\w.-]+:worksheet\b', xml):
        raise WorkflowError('Prefixed worksheet tags require a different preservation adapter.')
    pattern = cell_pattern(coordinate)
    found = pattern.search(xml)
    if found:
        return xml[:found.start()] + content_cell(found.group(), coordinate, value, formula, cached) + xml[found.end():]
    if cached:
        raise WorkflowError('Formula cell missing from original worksheet.')
    rownum, colnum = coordinate_to_tuple(coordinate)
    row = re.search(rb'<row\b[^>]*?\br="' + str(rownum).encode() + rb'"[^>]*?(?:/>|>.*?</row>)', xml, re.S)
    if not row or row.group().endswith(b'/>'):
        raise WorkflowError('Target row is not materialized in the template; structural expansion needs a separate explicit edit.')
    block = row.group()
    insertion = len(block) - len(b'</row>')
    for item in CELL.finditer(block):
        if coordinate_to_tuple(item.group(1).decode())[1] > colnum:
            insertion = item.start()
            break
    block = block[:insertion] + content_cell(None, coordinate, value, formula) + block[insertion:]
    return xml[:row.start()] + block + xml[row.end():]


def checked_edits(source, edits):
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter('always')
        wb = openpyxl.load_workbook(source, rich_text=True, keep_links=True)
    result, seen = [], set()
    for edit in edits:
        if not isinstance(edit, dict) or set(edit) not in ({'sheet', 'cell', 'value'}, {'sheet', 'cell', 'formula'}):
            raise WorkflowError('Each edit needs sheet, cell, and exactly one value or formula.')
        sheet, cell = edit['sheet'], edit['cell']
        if not isinstance(sheet, str) or sheet not in wb.sheetnames or not isinstance(cell, str) or not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]*', cell):
            raise WorkflowError('Unknown sheet or invalid single-cell coordinate.')
        if (sheet, cell) in seen:
            raise WorkflowError('Duplicate cell edits are ambiguous.')
        seen.add((sheet, cell))
        row, col = coordinate_to_tuple(cell)
        ws = wb[sheet]
        if row > ws.max_row or col > ws.max_column or isinstance(ws[cell], MergedCell) or ws.protection.sheet:
            raise WorkflowError('Edit outside template bounds, inside a non-anchor merged cell, or on a protected sheet.')
        if isinstance(ws[cell].value, CellRichText):
            raise WorkflowError('Replacing rich-text runs requires an explicit formatting-preserving adaptation.')
        is_formula = 'formula' in edit
        value = edit.get('formula') if is_formula else edit.get('value')
        if is_formula and (not isinstance(value, str) or not value.startswith('=') or len(value) < 2):
            raise WorkflowError('Formula edits must explicitly contain an equals-prefixed formula.')
        if value is not None and not isinstance(value, (str, bool, int, float)):
            raise WorkflowError('Only JSON scalar cell values are supported.')
        if isinstance(value, float) and not math.isfinite(value):
            raise WorkflowError('Non-finite numeric values are unsupported.')
        if isinstance(value, str) and (len(value) > 32767 or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', value)):
            raise WorkflowError('Cell text exceeds Excel bounds or contains XML control characters.')
        result.append((sheet, cell, value, is_formula))
    wb.close()
    return result, [str(x.message) for x in caught]


def edit_copy(source, destination, edits):
    source, destination = Path(source).resolve(), Path(destination).resolve()
    if source == destination or destination.exists():
        raise WorkflowError('Source and output must differ; existing outputs are not overwritten.')
    parts = package(source)
    if inspect(source)['blocked_features']:
        raise WorkflowError('Active content, external data, or signatures require another workflow; see inspect report.')
    validated, _ = checked_edits(source, edits)
    mapping = sheet_parts(parts)
    changed = set()
    replacements = {}
    for sheet, cell, value, is_formula in validated:
        part = mapping[sheet]
        original_cell = cell_pattern(cell).search(parts[part])
        if original_cell and re.search(rb'<f\b[^>]*\bt="shared"', original_cell.group()):
            raise WorkflowError('Editing a shared formula cell requires a group-aware adaptation.')
        replacements[part] = replace_cell(replacements.get(part, parts[part]), cell, value, is_formula)
        changed.add((sheet, cell))
    if not replacements:
        with source.open('rb') as src, destination.open('xb') as dst:
            shutil.copyfileobj(src, dst)
    else:
        write_package(source, destination, replacements)
    # Confirm exact intended semantics rather than merely ignoring allowed cells.
    wb = openpyxl.load_workbook(destination, rich_text=True)
    try:
        for sheet, cell, value, is_formula in validated:
            actual = wb[sheet][cell]
            actual_value = excel_scalar(actual.value, wb.epoch)
            if actual_value != value or (is_formula and actual.data_type != 'f') or (not is_formula and isinstance(value, str) and actual.data_type == 'f'):
                raise WorkflowError('Edited value/formula did not round-trip as requested.')
    finally:
        wb.close()
    return changed


def normalized_sheet(data, sheet, edited, cached):
    doc = ET.fromstring(data)
    for row in doc.findall(f'{{{MAIN}}}sheetData/{{{MAIN}}}row'):
        for cell in list(row):
            key = (sheet, cell.get('r'))
            if key in edited or key in cached:
                cell.attrib.pop('t', None)
                content = {'v', 'is', 'f'} if key in edited else {'v', 'is'}
                for child in list(cell):
                    if child.tag.rsplit('}', 1)[-1] in content:
                        cell.remove(child)
                if set(cell.attrib) == {'r'} and not len(cell):
                    row.remove(cell)
    return ET.canonicalize(ET.tostring(doc, encoding='unicode'))


def compare_packages(source, result, edited, cached):
    before, after = package(source), package(result)
    unexpected, altered = [], []
    mapping = {v: k for k, v in sheet_parts(before).items()}
    for name in sorted(set(before) | set(after)):
        if before.get(name) == after.get(name):
            continue
        altered.append(name)
        if name not in before or name not in after or name not in mapping:
            unexpected.append(name)
        elif normalized_sheet(before[name], mapping[name], edited, cached) != normalized_sheet(after[name], mapping[name], edited, cached):
            unexpected.append(name)
    return {'verified': not unexpected, 'altered_parts': altered, 'unexpected_changes': unexpected}


def run_tool(args, timeout):
    env = dict(os.environ, GSETTINGS_BACKEND='memory')
    process = subprocess.Popen(args, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                               env=env, start_new_session=True)
    try:
        stdout, stderr = process.communicate(timeout=timeout)
    except subprocess.TimeoutExpired:
        os.killpg(process.pid, signal.SIGKILL)
        process.communicate()
        raise WorkflowError('Tool deadline exceeded; process group terminated.')
    if process.returncode:
        raise WorkflowError(f'{Path(args[0]).name} failed with exit code {process.returncode}; no success artifact published.')
    return stdout.decode(errors='replace')


def office_convert(source, outdir, fmt, profile, timeout):
    outdir.mkdir()
    result = outdir / (source.stem + '.' + fmt)
    if fmt == 'xlsx':
        run_tool([sys.executable, str(Path(__file__).with_name('office_worker.py')),
                  str(source), str(result), str(profile)], timeout)
    else:
        run_tool([shutil.which('libreoffice'), '-env:UserInstallation=' + profile.as_uri(),
                  '--headless', '--nologo', '--nodefault', '--norestore', '--convert-to', fmt,
                  '--outdir', str(outdir), str(source)], timeout)
    if not result.is_file() or not result.stat().st_size:
        raise WorkflowError('LibreOffice returned no usable conversion artifact.')
    return result


def refresh_caches(source, destination, root, timeout):
    if inspect(source)['blocked_features']:
        raise WorkflowError('External-capable formulas or active workbook content cannot be recalculated by this workflow.')
    parts = package(source)
    temp_wb = ET.fromstring(parts['xl/workbook.xml'])
    calc = temp_wb.find(f'{{{MAIN}}}calcPr')
    if calc is None:
        calc = ET.SubElement(temp_wb, f'{{{MAIN}}}calcPr')
    calc.attrib.update(calcMode='auto', fullCalcOnLoad='1', forceFullCalc='1')
    recalc_input = root / 'recalculate.xlsx'
    write_package(source, recalc_input, {'xl/workbook.xml': ET.tostring(temp_wb)})
    evaluated = office_convert(recalc_input, root / 'recalculated', 'xlsx', root / 'office-profile', timeout)
    formulas = openpyxl.load_workbook(source, data_only=False)
    values = openpyxl.load_workbook(evaluated, data_only=True)
    evaluated_parts = package(evaluated)
    blank_strings = set()
    for sheet, part in sheet_parts(evaluated_parts).items():
        for node in ET.fromstring(evaluated_parts[part]).iter(f'{{{MAIN}}}c'):
            payload = node.find(f'{{{MAIN}}}v')
            if node.get('t') == 'str' and payload is not None and not payload.text:
                blank_strings.add((sheet, node.get('r')))
    mapping = sheet_parts(parts)
    replacements, cached, errors = {}, set(), []
    if formulas.sheetnames != values.sheetnames:
        raise WorkflowError('Recalculation changed sheet identity/order.')
    for ws in formulas:
        for row in ws:
            for cell in row:
                if cell.data_type != 'f':
                    continue
                if not isinstance(cell.value, str):
                    raise WorkflowError('Array/shared formula objects require a specialized cache adapter.')
                evaluated_cell = values[ws.title][cell.coordinate]
                value = evaluated_cell.value
                if evaluated_cell.data_type == 'e' or (value is None and (ws.title, cell.coordinate) not in blank_strings):
                    errors.append({'sheet': ws.title, 'cell': cell.coordinate, 'code': 'formula_error_or_missing_cache'})
                    continue
                if value is None:
                    value = ''
                value = excel_scalar(value, formulas.epoch)
                part = mapping[ws.title]
                replacements[part] = replace_cell(replacements.get(part, parts[part]), cell.coordinate, value, cached=True)
                cached.add((ws.title, cell.coordinate))
    formulas.close()
    values.close()
    if errors:
        raise WorkflowError('Formula errors or missing caches at ' + ', '.join(x['sheet'] + '!' + x['cell'] for x in errors))
    write_package(source, destination, replacements)
    return cached


def runtime_status():
    result = {'tools': {name: shutil.which(name) for name in ('libreoffice', 'pdftoppm', 'pdfinfo')},
              'local_socket_bind': False, 'python_uno_available': importlib.util.find_spec('uno') is not None}
    with tempfile.TemporaryDirectory(prefix='spreadsheet-runtime-check-') as folder:
        handle = None
        try:
            handle = socket.socket(socket.AF_UNIX, socket.SOCK_STREAM)
            handle.bind(str(Path(folder) / 'check.sock'))
            result['local_socket_bind'] = True
        except OSError as error:
            result['socket_error'] = type(error).__name__ + ': errno ' + str(error.errno)
        finally:
            if handle is not None:
                handle.close()
    result['ready_for_office_probe'] = all(result['tools'].values()) and result['local_socket_bind'] and result['python_uno_available']
    # Installed executables and a socket check are prerequisites, not a successful conversion.
    return result


def build(source, output_dir, edits, timeout=60, max_pages=20, dpi=120):
    source, output_dir = Path(source).resolve(), Path(output_dir).resolve()
    if timeout <= 0 or max_pages < 1 or not 40 <= dpi <= 300:
        raise WorkflowError('Require positive timeout/pages and DPI between 40 and 300.')
    if output_dir.exists():
        raise WorkflowError('Choose a new output directory; existing artifacts are never overwritten.')
    initial = inspect(source)
    output_dir.mkdir(parents=True)
    report = dict(initial, verified=False, outputs=[], visual_review='pending')
    try:
        if initial['blocked_features']:
            raise WorkflowError('Blocked active/external/signed workbook features; source is unchanged.')
        for tool in ('libreoffice', 'pdftoppm', 'pdfinfo'):
            if not shutil.which(tool):
                raise WorkflowError('Missing local tool: ' + tool)
        with tempfile.TemporaryDirectory(prefix='codex-spreadsheets-') as temp:
            root = Path(temp)
            edited = root / 'edited.xlsx'
            changed = edit_copy(source, edited, edits)
            final = root / 'workbook.xlsx'
            cached = refresh_caches(edited, final, root, timeout)
            fidelity = compare_packages(source, final, changed, cached)
            report.update(fidelity)
            if not fidelity['verified']:
                raise WorkflowError('Unrequested package changes detected; inspect unexpected_changes.')
            report['verified'] = False  # Success additionally requires actual preview artifacts.
            pdf = office_convert(final, root / 'pdf', 'pdf', root / 'pdf-profile', timeout)
            info = run_tool([shutil.which('pdfinfo'), str(pdf)], timeout)
            match = re.search(r'^Pages:\s*(\d+)', info, re.M)
            if not match or not 1 <= int(match.group(1)) <= max_pages:
                raise WorkflowError('Preview exceeds page limit or page count is unavailable; adjust print layout explicitly or raise --max-pages.')
            run_tool([shutil.which('pdftoppm'), '-r', str(dpi), '-png', str(pdf), str(root / 'page')], timeout)
            previews = sorted(root.glob('page-*.png'))
            if len(previews) != int(match.group(1)):
                raise WorkflowError('Not all preview pages were rendered.')
            from PIL import Image
            for preview in previews:
                with Image.open(preview) as image:
                    image.verify()
            if sha(source) != initial['source_sha256']:
                raise WorkflowError('Source changed during workflow; refusing success.')
            for src, name in [(final, 'workbook.xlsx'), (pdf, 'preview.pdf')] + [(p, p.name) for p in previews]:
                with src.open('rb') as inp, (output_dir / name).open('xb') as out:
                    shutil.copyfileobj(inp, out)
                report['outputs'].append(name)
            report.update(verified=True, edited_cells=len(changed), formula_cells_recalculated=len(cached),
                          pages=len(previews), source_unchanged=True,
                          verification_scope='Package fidelity and preview readability; agent visual review and task-specific assertions still required.')
    except Exception as error:
        report['verified'] = False
        report['error'] = str(error) if isinstance(error, WorkflowError) else type(error).__name__
        # Do not leave a partial final workbook advertised after a failed build.
        for name in report['outputs']:
            (output_dir / name).unlink(missing_ok=True)
        report['outputs'] = []
        raise WorkflowError(report['error']) from error
    finally:
        report['source_unchanged'] = sha(source) == initial['source_sha256']
        (output_dir / 'report.json').write_text(json.dumps(report, indent=2))
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest='command', required=True)
    inspect_cmd = commands.add_parser('inspect')
    inspect_cmd.add_argument('source', type=Path)
    commands.add_parser('doctor')
    create = commands.add_parser('build')
    create.add_argument('source', type=Path)
    create.add_argument('--edits', type=Path)
    create.add_argument('--output-dir', type=Path, required=True)
    create.add_argument('--timeout', type=int, default=60)
    create.add_argument('--max-pages', type=int, default=20)
    create.add_argument('--dpi', type=int, default=120)
    args = parser.parse_args()
    try:
        if args.command == 'doctor':
            result = runtime_status()
        elif args.command == 'inspect':
            result = inspect(args.source)
        else:
            edits = []
            if args.edits:
                payload = json.loads(args.edits.read_text())
                if not isinstance(payload, dict) or set(payload) != {'edits'} or not isinstance(payload['edits'], list):
                    raise WorkflowError('Edit file must contain only an edits list.')
                edits = payload['edits']
            result = build(args.source, args.output_dir, edits, args.timeout, args.max_pages, args.dpi)
        print(json.dumps(result, indent=2))
    except (WorkflowError, ValueError, OSError) as error:
        parser.exit(1, f'{error}\n')


if __name__ == '__main__':
    main()
