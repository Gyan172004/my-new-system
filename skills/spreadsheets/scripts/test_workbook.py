"""Behavioral checks; all artifacts stay in temporary directories."""
import hashlib
import importlib.util
import json
import os
import shutil
import sys
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.defined_name import DefinedName

spec = importlib.util.spec_from_file_location('workbook', Path(__file__).with_name('workbook.py'))
workbook = importlib.util.module_from_spec(spec)
spec.loader.exec_module(workbook)


class WorkbookTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.source = self.root / 'source.xlsx'
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Example'
        ws.append(['Quantity', 'Price', 'Total', 'Note'])
        ws.append([2, 7, '=A2*B2', 'Original'])
        ws.append([1, 4, '=A3*B3', ''])
        ws['C2'].number_format = '$0.00'
        ws['A1'].font = Font(bold=True, color='FFFFFF')
        ws['A1'].fill = PatternFill('solid', fgColor='17365D')
        ws.freeze_panes = 'A2'
        ws.column_dimensions['D'].width = 28
        ws.row_dimensions[1].height = 26
        ws.merge_cells('A5:D5')
        ws['A5'] = 'Local workflow demonstration'
        ws.print_area = 'A1:D5'
        table = Table(displayName='Items', ref='A1:D3')
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        ws.add_table(table)
        dv = DataValidation(type='whole', operator='greaterThan', formula1=0)
        dv.add('A2:A3')
        ws.add_data_validation(dv)
        ws.conditional_formatting.add('C2:C3', CellIsRule(operator='greaterThan', formula=['10'], font=Font(bold=True)))
        wb.defined_names.add(DefinedName('UnitPrice', attr_text="'Example'!$B$2"))
        wb.save(self.source)
        self.original = hashlib.sha256(self.source.read_bytes()).hexdigest()

    def tearDown(self):
        self.tmp.cleanup()

    def edits(self, *items):
        p = self.root / 'edits.json'
        p.write_text(json.dumps({'edits': list(items)}))
        return p

    def test_literal_edits_preserve_template_and_package_features(self):
        destination = self.root / 'edited.xlsx'
        changes = workbook.edit_copy(self.source, destination, [
            {'sheet': 'Example', 'cell': 'A2', 'value': 3},
            {'sheet': 'Example', 'cell': 'D2', 'value': '=literal text'},
        ])
        wb = openpyxl.load_workbook(destination)
        self.assertEqual(wb['Example']['A2'].value, 3)
        self.assertEqual(wb['Example']['D2'].value, '=literal text')
        self.assertEqual(wb['Example']['D2'].data_type, 's')
        self.assertEqual(wb['Example']['C2'].value, '=A2*B2')
        self.assertEqual(wb['Example'].freeze_panes, 'A2')
        self.assertIn('Items', wb['Example'].tables)
        self.assertTrue(workbook.compare_packages(self.source, destination, changes, set())['verified'])
        self.assertEqual(hashlib.sha256(self.source.read_bytes()).hexdigest(), self.original)

    def test_clone_is_byte_identical_and_refuses_overwrite(self):
        destination = self.root / 'clone.xlsx'
        workbook.edit_copy(self.source, destination, [])
        self.assertEqual(destination.read_bytes(), self.source.read_bytes())
        with self.assertRaises(workbook.WorkflowError):
            workbook.edit_copy(self.source, self.source, [])
        with self.assertRaises(workbook.WorkflowError):
            workbook.edit_copy(self.source, destination, [])

    def test_edits_are_explicit_bounded_and_do_not_execute_formula_like_text(self):
        invalid = [
            [{'sheet': 'Example', 'cell': 'A5:B5', 'value': 'bad'}],
            [{'sheet': 'Example', 'cell': 'B5', 'value': 'merged'}],
            [{'sheet': 'Example', 'cell': 'A100', 'value': 'outside'}],
            [{'sheet': 'Example', 'cell': 'A2', 'value': 1, 'formula': '=2'}],
            [{'sheet': 'Example', 'cell': 'A2', 'value': float('nan')}],
            [{'sheet': 'Example', 'cell': 'A2', 'formula': '2+2'}],
            [{'sheet': 'Example', 'cell': 'A2', 'value': 1}] * 2,
        ]
        for edits in invalid:
            with self.subTest(edits=edits):
                with self.assertRaises(workbook.WorkflowError):
                    workbook.edit_copy(self.source, self.root / 'rejected.xlsx', edits)
        self.assertFalse((self.root / 'rejected.xlsx').exists())

    def test_fidelity_check_detects_unrequested_style_change(self):
        changed = self.root / 'changed.xlsx'
        wb = openpyxl.load_workbook(self.source)
        wb['Example']['A1'].font = Font(bold=False)
        wb.save(changed)
        report = workbook.compare_packages(self.source, changed, set(), set())
        self.assertFalse(report['verified'])
        self.assertTrue(report['unexpected_changes'])

    def test_shared_formula_edit_requires_group_adaptation(self):
        shared = self.root / 'shared.xlsx'
        parts = workbook.package(self.source)
        part = workbook.sheet_parts(parts)['Example']
        xml = parts[part].replace(b'<f>A2*B2</f>', b'<f t="shared" si="0" ref="C2:C3">A2*B2</f>')
        xml = xml.replace(b'<f>A3*B3</f>', b'<f t="shared" si="0"/>')
        workbook.write_package(self.source, shared, {part: xml})
        self.assertEqual(openpyxl.load_workbook(shared)['Example']['C3'].value, '=A3*B3')
        destination = self.root / 'broken-group.xlsx'
        with self.assertRaisesRegex(workbook.WorkflowError, 'shared formula'):
            workbook.edit_copy(shared, destination, [{'sheet': 'Example', 'cell': 'C2', 'value': 42}])
        self.assertFalse(destination.exists())

    def test_empty_cell_edit_does_not_consume_following_formula_cells(self):
        source = self.root / 'self-closing-cell.xlsx'
        parts = workbook.package(self.source)
        part = workbook.sheet_parts(parts)['Example']
        original = workbook.cell_pattern('B3').search(parts[part]).group()
        xml = parts[part].replace(original, b'<c r="B3" s="0"/>')
        workbook.write_package(self.source, source, {part: xml})
        destination = self.root / 'edited-empty.xlsx'
        changes = workbook.edit_copy(source, destination, [{'sheet': 'Example', 'cell': 'B3', 'value': 8}])
        wb = openpyxl.load_workbook(destination)
        self.assertEqual(wb['Example']['B3'].value, 8)
        self.assertEqual(wb['Example']['C3'].value, '=A3*B3')
        self.assertTrue(workbook.compare_packages(source, destination, changes, set())['verified'])

    def test_scalar_text_escaping_booleans_numbers_and_clear(self):
        dest = self.root / 'scalars.xlsx'
        changes = workbook.edit_copy(self.source, dest, [
            {'sheet': 'Example', 'cell': 'A2', 'value': 3.25},
            {'sheet': 'Example', 'cell': 'B2', 'value': False},
            {'sheet': 'Example', 'cell': 'D2', 'value': '<tag> & "literal"'},
            {'sheet': 'Example', 'cell': 'D3', 'value': None},
        ])
        wb = openpyxl.load_workbook(dest)
        self.assertEqual(wb['Example']['A2'].value, 3.25)
        self.assertIs(wb['Example']['B2'].value, False)
        self.assertEqual(wb['Example']['D2'].value, '<tag> & "literal"')
        self.assertIsNone(wb['Example']['D3'].value)
        self.assertTrue(workbook.compare_packages(self.source, dest, changes, set())['verified'])

    def test_numeric_excel_date_edit_retains_date_format_and_meaning(self):
        from datetime import datetime
        wb = openpyxl.load_workbook(self.source)
        wb['Example']['B2'].number_format = 'mmm d, yyyy'
        dated = self.root / 'date-template.xlsx'
        wb.save(dated)
        destination = self.root / 'date-edited.xlsx'
        value = openpyxl.utils.datetime.to_excel(datetime(2026, 10, 17))
        changes = workbook.edit_copy(dated, destination, [{'sheet': 'Example', 'cell': 'B2', 'value': value}])
        actual = openpyxl.load_workbook(destination)['Example']['B2']
        self.assertEqual(actual.value, datetime(2026, 10, 17))
        self.assertEqual(actual.number_format, 'mmm d, yyyy')
        self.assertTrue(workbook.compare_packages(dated, destination, changes, set())['verified'])

    def test_tool_deadline_stops_process(self):
        with self.assertRaisesRegex(workbook.WorkflowError, 'deadline exceeded'):
            workbook.run_tool([sys.executable, '-c', 'import time; time.sleep(10)'], timeout=0.1)

    @unittest.skipUnless(shutil.which('pdftoppm') and shutil.which('pdfinfo'), 'Poppler tools unavailable.')
    def test_real_pdf_rasterization_without_claiming_office_conversion(self):
        from PIL import Image
        pdf = self.root / 'synthetic.pdf'
        Image.new('RGB', (120, 80), 'white').save(pdf, 'PDF')
        workbook.run_tool([shutil.which('pdftoppm'), '-r', '72', '-png', str(pdf), str(self.root / 'raster')], timeout=10)
        pngs = list(self.root.glob('raster-*.png'))
        self.assertEqual(len(pngs), 1)
        with Image.open(pngs[0]) as image:
            image.verify()

    @unittest.skipUnless(os.environ.get('RUN_OFFICE_TESTS') == '1', 'Office integration requires a runtime permitting local sockets; enable with RUN_OFFICE_TESTS=1.')
    def test_formula_recalculation_and_rendered_artifacts(self):
        out = self.root / 'result'
        report = workbook.build(self.source, out, [
            {'sheet': 'Example', 'cell': 'A2', 'value': 3},
            {'sheet': 'Example', 'cell': 'D3', 'formula': '=SUM(C2:C3)'},
        ], timeout=45, max_pages=5, dpi=80)
        self.assertTrue(report['verified'])
        cached = openpyxl.load_workbook(out / 'workbook.xlsx', data_only=True)
        formulas = openpyxl.load_workbook(out / 'workbook.xlsx', data_only=False)
        self.assertEqual(cached['Example']['C2'].value, 21)
        self.assertEqual(cached['Example']['D3'].value, 25)
        self.assertEqual(formulas['Example']['D3'].value, '=SUM(C2:C3)')
        self.assertTrue((out / 'preview.pdf').is_file())
        self.assertTrue(list(out.glob('page-*.png')))
        self.assertEqual(hashlib.sha256(self.source.read_bytes()).hexdigest(), self.original)

    @unittest.skipUnless(os.environ.get('RUN_OFFICE_TESTS') == '1', 'Office integration requires a runtime permitting local sockets; enable with RUN_OFFICE_TESTS=1.')
    def test_formula_error_is_reported_without_publishing_success(self):
        out = self.root / 'failed'
        with self.assertRaises(workbook.WorkflowError):
            workbook.build(self.source, out, [{'sheet': 'Example', 'cell': 'C2', 'formula': '=1/0'}], timeout=45, max_pages=5, dpi=80)
        report = json.loads((out / 'report.json').read_text())
        self.assertFalse(report['verified'])
        self.assertFalse((out / 'workbook.xlsx').exists())

    def test_cache_transplant_preserves_formulas_styles_and_other_zip_parts(self):
        evaluated = self.root / 'evaluated.xlsx'
        parts = workbook.package(self.source)
        part = workbook.sheet_parts(parts)['Example']
        xml = workbook.replace_cell(parts[part], 'C2', 14, cached=True)
        xml = workbook.replace_cell(xml, 'C3', 4, cached=True)
        workbook.write_package(self.source, evaluated, {part: xml})
        final = self.root / 'final.xlsx'
        # This unit test verifies cache transplant mechanics; it does not run Office.
        with patch.object(workbook, 'office_convert', return_value=evaluated):
            cached = workbook.refresh_caches(self.source, final, self.root, timeout=30)
        self.assertEqual(openpyxl.load_workbook(final, data_only=True)['Example']['C2'].value, 14)
        self.assertEqual(openpyxl.load_workbook(final)['Example']['C2'].value, '=A2*B2')
        self.assertTrue(workbook.compare_packages(self.source, final, set(), cached)['verified'])
        with ZipFile(self.source) as before, ZipFile(final) as after:
            for name in before.namelist():
                if name != part:
                    self.assertEqual(before.read(name), after.read(name), name)

    def test_empty_string_formula_cache_is_valid_but_missing_cache_is_not(self):
        parts = workbook.package(self.source)
        part = workbook.sheet_parts(parts)['Example']
        for present in [True, False]:
            with self.subTest(cache_present=present):
                evaluated = self.root / ('blank-cache-' + str(present) + '.xlsx')
                xml = workbook.replace_cell(parts[part], 'C2', '', cached=True)
                xml = workbook.replace_cell(xml, 'C3', 4, cached=True)
                if not present:
                    xml = xml.replace(b'<f>A2*B2</f><v></v>', b'<f>A2*B2</f>')
                workbook.write_package(self.source, evaluated, {part: xml})
                stage = self.root / ('stage-' + str(present))
                stage.mkdir()
                final = stage / 'final.xlsx'
                with patch.object(workbook, 'office_convert', return_value=evaluated):
                    if present:
                        workbook.refresh_caches(self.source, final, stage, timeout=30)
                        self.assertEqual(openpyxl.load_workbook(final, data_only=True)['Example']['C2'].data_type, 'str')
                    else:
                        with self.assertRaises(workbook.WorkflowError):
                            workbook.refresh_caches(self.source, final, stage, timeout=30)

    @unittest.skipUnless(os.environ.get('RUN_OFFICE_TESTS') == '1', 'Office integration requires RUN_OFFICE_TESTS=1.')
    def test_recalculation_replaces_existing_stale_formula_caches(self):
        source = self.root / 'stale-caches.xlsx'
        parts = workbook.package(self.source)
        part = workbook.sheet_parts(parts)['Example']
        xml = workbook.replace_cell(parts[part], 'C2', 999, cached=True)
        xml = workbook.replace_cell(xml, 'C3', 888, cached=True)
        workbook.write_package(self.source, source, {part: xml})
        report = workbook.build(source, self.root / 'fresh-result', [{'sheet': 'Example', 'cell': 'A2', 'value': 3}], timeout=45)
        self.assertTrue(report['verified'])
        ws = openpyxl.load_workbook(self.root / 'fresh-result/workbook.xlsx', data_only=True)['Example']
        self.assertEqual(ws['C2'].value, 21)
        self.assertEqual(ws['C3'].value, 4)

    def test_renderer_failure_retains_failure_report_and_no_final_workbook(self):
        out = self.root / 'renderer-failure'
        with patch.object(workbook, 'office_convert', side_effect=workbook.WorkflowError('conversion unavailable')):
            with self.assertRaises(workbook.WorkflowError):
                workbook.build(self.source, out, [], timeout=10, max_pages=2, dpi=80)
        report = json.loads((out / 'report.json').read_text())
        self.assertFalse(report['verified'])
        self.assertTrue(report['source_unchanged'])
        self.assertEqual(report['outputs'], [])
        self.assertFalse((out / 'workbook.xlsx').exists())

    def test_network_formula_is_blocked_before_conversion(self):
        dest = self.root / 'external-formula.xlsx'
        workbook.edit_copy(self.source, dest, [{'sheet': 'Example', 'cell': 'C2', 'formula': '=WEBSERVICE("https://example.invalid")'}])
        self.assertTrue(workbook.inspect(dest)['blocked_features'])
        with patch.object(workbook, 'office_convert') as converter:
            with self.assertRaises(workbook.WorkflowError):
                workbook.build(dest, self.root / 'no-network', [])
        converter.assert_not_called()

    def test_unsafe_active_content_is_reported(self):
        special = self.root / 'active.xlsx'
        with ZipFile(self.source) as src, ZipFile(special, 'w', ZIP_DEFLATED) as dst:
            for info in src.infolist():
                dst.writestr(info, src.read(info.filename))
            dst.writestr('xl/externalLinks/externalLink1.xml', '<externalLink/>')
        info = workbook.inspect(special)
        self.assertTrue(info['blocked_features'])
        with self.assertRaises(workbook.WorkflowError):
            workbook.build(special, self.root / 'blocked', [], timeout=45, max_pages=5, dpi=80)


if __name__ == '__main__':
    unittest.main()
